from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout,
    QPushButton, QLabel, QFrame,
    QTableWidget, QTableWidgetItem,
    QHeaderView, QDateEdit, QComboBox,
    QMessageBox, QFileDialog, QTabWidget
)
from PyQt6.QtCore import Qt, QDate
from PyQt6.QtGui import QColor
from models.sale import Sale
from models.product import Product
from database.db import db
import os

class ReportsView(QWidget):

    def __init__(self):
        super().__init__()
        self.build_ui()

    def build_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(30, 25, 30, 25)
        layout.setSpacing(15)

        # Header
        header = QHBoxLayout()
        left = QVBoxLayout()
        title = QLabel("📊 Reports")
        title.setObjectName("page_title")
        subtitle = QLabel("Sales analytics and business insights")
        subtitle.setObjectName("page_subtitle")
        left.addWidget(title)
        left.addWidget(subtitle)
        header.addLayout(left)
        header.addStretch()

        export_btn = QPushButton("📥 Export Excel")
        export_btn.setStyleSheet("""
            QPushButton {
                background-color: #00c853;
                color: white;
                border-radius: 8px;
                padding: 10px 20px;
                font-weight: bold;
            }
            QPushButton:hover { background-color: #00e676; }
        """)
        export_btn.clicked.connect(self.export_excel)
        header.addWidget(export_btn)
        layout.addLayout(header)

        # Date filter
        filter_frame = QFrame()
        filter_frame.setStyleSheet("""
            QFrame {
                background-color: #1a1a2e;
                border: 1px solid #2a2a4a;
                border-radius: 10px;
                padding: 10px;
            }
        """)
        filter_layout = QHBoxLayout(filter_frame)
        filter_layout.setContentsMargins(15, 10, 15, 10)

        period_label = QLabel("Period:")
        period_label.setStyleSheet("color: #6a6a8a;")
        self.period_combo = QComboBox()
        self.period_combo.addItems([
            "Today", "Yesterday", "This Week",
            "This Month", "Last Month", "Custom Range"
        ])
        self.period_combo.currentIndexChanged.connect(
            self.on_period_changed)
        self.period_combo.setStyleSheet("""
            QComboBox {
                background-color: #0f0f1a;
                color: #e0e0e0;
                border: 1px solid #2a2a4a;
                border-radius: 6px;
                padding: 6px 12px;
                min-width: 150px;
            }
        """)

        from_label = QLabel("From:")
        from_label.setStyleSheet("color: #6a6a8a;")
        self.from_date = QDateEdit()
        self.from_date.setDate(QDate.currentDate())
        self.from_date.setCalendarPopup(True)
        self.from_date.setStyleSheet("""
            QDateEdit {
                background-color: #0f0f1a;
                color: #e0e0e0;
                border: 1px solid #2a2a4a;
                border-radius: 6px;
                padding: 6px 12px;
            }
        """)

        to_label = QLabel("To:")
        to_label.setStyleSheet("color: #6a6a8a;")
        self.to_date = QDateEdit()
        self.to_date.setDate(QDate.currentDate())
        self.to_date.setCalendarPopup(True)
        self.to_date.setStyleSheet(self.from_date.styleSheet())

        apply_btn = QPushButton("Apply")
        apply_btn.setStyleSheet("""
            QPushButton {
                background-color: #4a9eff;
                color: white;
                border-radius: 6px;
                padding: 8px 20px;
                font-weight: bold;
            }
        """)
        apply_btn.clicked.connect(self.load_report)

        filter_layout.addWidget(period_label)
        filter_layout.addWidget(self.period_combo)
        filter_layout.addWidget(from_label)
        filter_layout.addWidget(self.from_date)
        filter_layout.addWidget(to_label)
        filter_layout.addWidget(self.to_date)
        filter_layout.addWidget(apply_btn)
        filter_layout.addStretch()
        layout.addWidget(filter_frame)

        # Summary cards
        self.cards_layout = QHBoxLayout()
        self.cards_layout.setSpacing(12)
        self.card_revenue = self.make_card(
            "💰 Revenue", "GHS 0.00", "#4a9eff")
        self.card_profit = self.make_card(
            "📈 Profit", "GHS 0.00", "#00c853")
        self.card_transactions = self.make_card(
            "🛒 Sales", "0", "#ffab40")
        self.card_avg = self.make_card(
            "📊 Avg Sale", "GHS 0.00", "#cba6f7")
        self.card_cash = self.make_card(
            "💵 Cash", "GHS 0.00", "#4a9eff")
        self.card_momo = self.make_card(
            "📱 MoMo", "GHS 0.00", "#00c853")

        for card in [self.card_revenue, self.card_profit,
                     self.card_transactions, self.card_avg,
                     self.card_cash, self.card_momo]:
            self.cards_layout.addWidget(card)
        layout.addLayout(self.cards_layout)

        # Tabs
        tabs = QTabWidget()
        tabs.setStyleSheet("""
            QTabWidget::pane {
                border: 1px solid #2a2a4a;
                border-radius: 8px;
                background-color: #1a1a2e;
            }
            QTabBar::tab {
                background-color: #1a1a2e;
                color: #6a6a8a;
                padding: 10px 20px;
                border: none;
                border-bottom: 2px solid transparent;
            }
            QTabBar::tab:selected {
                color: #4a9eff;
                border-bottom: 2px solid #4a9eff;
            }
        """)

        # Sales tab
        sales_tab = QWidget()
        sales_layout = QVBoxLayout(sales_tab)
        sales_layout.setContentsMargins(10, 10, 10, 10)
        self.sales_table = self.make_sales_table()
        sales_layout.addWidget(self.sales_table)
        tabs.addTab(sales_tab, "Sales Transactions")

        # Top products tab
        products_tab = QWidget()
        products_layout = QVBoxLayout(products_tab)
        products_layout.setContentsMargins(10, 10, 10, 10)
        self.products_table = self.make_products_table()
        products_layout.addWidget(self.products_table)
        tabs.addTab(products_tab, "🏆 Top Products")

        layout.addWidget(tabs)
        self.load_report()

    def make_card(self, label, value, color):
        frame = QFrame()
        frame.setObjectName("stat_card")
        frame.setSizePolicy(
            frame.sizePolicy().horizontalPolicy(),
            frame.sizePolicy().verticalPolicy()
        )
        layout = QVBoxLayout(frame)
        layout.setContentsMargins(12, 12, 12, 12)
        lbl = QLabel(label)
        lbl.setStyleSheet("color: #6a6a8a; font-size: 11px;")
        val = QLabel(value)
        val.setStyleSheet(
            f"color: {color}; font-size: 18px; font-weight: bold;")
        layout.addWidget(lbl)
        layout.addWidget(val)
        frame.value_label = val
        return frame

    def make_sales_table(self):
        table = QTableWidget()
        table.setColumnCount(7)
        table.setHorizontalHeaderLabels([
            "ID", "Customer", "Subtotal",
            "Discount", "Total", "Payment", "Date/Time"
        ])
        table.horizontalHeader().setSectionResizeMode(
            1, QHeaderView.ResizeMode.Stretch)
        table.setColumnWidth(0, 50)
        table.setColumnWidth(2, 100)
        table.setColumnWidth(3, 90)
        table.setColumnWidth(4, 100)
        table.setColumnWidth(5, 90)
        table.setColumnWidth(6, 150)
        table.verticalHeader().setVisible(False)
        table.setAlternatingRowColors(True)
        table.setEditTriggers(
            QTableWidget.EditTrigger.NoEditTriggers)
        return table

    def make_products_table(self):
        table = QTableWidget()
        table.setColumnCount(3)
        table.setHorizontalHeaderLabels([
            "Product", "Units Sold", "Revenue"
        ])
        table.horizontalHeader().setSectionResizeMode(
            0, QHeaderView.ResizeMode.Stretch)
        table.setColumnWidth(1, 120)
        table.setColumnWidth(2, 150)
        table.verticalHeader().setVisible(False)
        table.setAlternatingRowColors(True)
        table.setEditTriggers(
            QTableWidget.EditTrigger.NoEditTriggers)
        return table

    def on_period_changed(self, index):
        today = QDate.currentDate()
        if index == 0:
            self.from_date.setDate(today)
            self.to_date.setDate(today)
        elif index == 1:
            yesterday = today.addDays(-1)
            self.from_date.setDate(yesterday)
            self.to_date.setDate(yesterday)
        elif index == 2:
            self.from_date.setDate(
                today.addDays(-today.dayOfWeek() + 1))
            self.to_date.setDate(today)
        elif index == 3:
            self.from_date.setDate(
                QDate(today.year(), today.month(), 1))
            self.to_date.setDate(today)
        elif index == 4:
            last = today.addMonths(-1)
            self.from_date.setDate(
                QDate(last.year(), last.month(), 1))
            self.to_date.setDate(
                QDate(last.year(), last.month(),
                      last.daysInMonth()))
        self.load_report()

    def load_report(self):
        start = self.from_date.date().toString("yyyy-MM-dd")
        end = self.to_date.date().toString("yyyy-MM-dd")

        sales = Sale.get_sales_by_date(start, end)
        profit = Sale.get_profit_summary(start, end)

        revenue = profit["total_revenue"] or 0
        profit_val = profit["total_profit"] or 0
        total_sales = profit["total_sales"] or 0
        avg_sale = (revenue / total_sales
                    if total_sales > 0 else 0)

        cash_total = sum(
            s["total_amount"] for s in sales
            if s["payment_method"] == "cash")
        momo_total = sum(
            s["total_amount"] for s in sales
            if s["payment_method"] == "momo")

        self.card_revenue.value_label.setText(
            "GHS {:.2f}".format(revenue))
        self.card_profit.value_label.setText(
            "GHS {:.2f}".format(profit_val))
        self.card_transactions.value_label.setText(
            str(total_sales))
        self.card_avg.value_label.setText(
            "GHS {:.2f}".format(avg_sale))
        self.card_cash.value_label.setText(
            "GHS {:.2f}".format(cash_total))
        self.card_momo.value_label.setText(
            "GHS {:.2f}".format(momo_total))

        self.sales_table.setRowCount(len(sales))
        for row, s in enumerate(sales):
            self.sales_table.setItem(
                row, 0, QTableWidgetItem(str(s["id"])))
            self.sales_table.setItem(
                row, 1, QTableWidgetItem(
                    s.get("customer_name") or "Walk-in"))
            self.sales_table.setItem(
                row, 2, QTableWidgetItem(
                    "GHS {:.2f}".format(s["subtotal"])))
            self.sales_table.setItem(
                row, 3, QTableWidgetItem(
                    "GHS {:.2f}".format(s["discount_amount"])))
            self.sales_table.setItem(
                row, 4, QTableWidgetItem(
                    "GHS {:.2f}".format(s["total_amount"])))
            pay_item = QTableWidgetItem(
                s["payment_method"].upper())
            if s["payment_method"] == "momo":
                pay_item.setForeground(QColor("#00c853"))
            elif s["payment_method"] == "credit":
                pay_item.setForeground(QColor("#ff1744"))
            self.sales_table.setItem(row, 5, pay_item)
            self.sales_table.setItem(
                row, 6, QTableWidgetItem(
                    s["created_at"][:16]))

        top_products = Sale.get_top_products(limit=20)
        self.products_table.setRowCount(len(top_products))
        for row, p in enumerate(top_products):
            self.products_table.setItem(
                row, 0, QTableWidgetItem(p["product_name"]))
            sold_item = QTableWidgetItem(str(p["total_sold"]))
            sold_item.setTextAlignment(
                Qt.AlignmentFlag.AlignCenter)
            self.products_table.setItem(row, 1, sold_item)
            self.products_table.setItem(
                row, 2, QTableWidgetItem(
                    "GHS {:.2f}".format(p["total_revenue"])))

    def export_excel(self):
        try:
            import openpyxl
            from openpyxl.styles import (
                PatternFill, Font, Alignment)

            start = self.from_date.date().toString("yyyy-MM-dd")
            end = self.to_date.date().toString("yyyy-MM-dd")

            path, _ = QFileDialog.getSaveFileName(
                self, "Export Report",
                f"SmartPOS_Report_{start}_to_{end}.xlsx",
                "Excel Files (*.xlsx)"
            )
            if not path:
                return

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Sales Report"

            header_fill = PatternFill(
                "solid", fgColor="1a1a2e")
            header_font = Font(
                bold=True, color="4a9eff")

            headers = [
                "Sale ID", "Customer", "Subtotal",
                "Discount", "Total", "Payment", "Date"
            ]
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=h)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")

            sales = Sale.get_sales_by_date(start, end)
            for row, s in enumerate(sales, 2):
                ws.cell(row=row, column=1, value=s["id"])
                ws.cell(row=row, column=2,
                        value=s.get("customer_name") or "Walk-in")
                ws.cell(row=row, column=3,
                        value=round(s["subtotal"], 2))
                ws.cell(row=row, column=4,
                        value=round(s["discount_amount"], 2))
                ws.cell(row=row, column=5,
                        value=round(s["total_amount"], 2))
                ws.cell(row=row, column=6,
                        value=s["payment_method"].upper())
                ws.cell(row=row, column=7,
                        value=s["created_at"][:16])

            profit = Sale.get_profit_summary(start, end)
            ws2 = wb.create_sheet("Summary")
            ws2["A1"] = "SmartPOS Ghana — Sales Summary"
            ws2["A1"].font = Font(bold=True, size=14)
            ws2["A3"] = "Period:"
            ws2["B3"] = f"{start} to {end}"
            ws2["A4"] = "Total Revenue:"
            ws2["B4"] = round(profit["total_revenue"] or 0, 2)
            ws2["A5"] = "Total Profit:"
            ws2["B5"] = round(profit["total_profit"] or 0, 2)
            ws2["A6"] = "Total Sales:"
            ws2["B6"] = profit["total_sales"] or 0

            wb.save(path)
            QMessageBox.information(
                self, "Export Complete",
                f"Report saved to:\n{path}")

        except Exception as e:
            QMessageBox.critical(
                self, "Export Failed", str(e))

    def refresh(self):
        self.load_report()